from __future__ import annotations

import io
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape

import fitz
from docx import Document
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def inspect_docx(raw: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(item.file_size for item in entries) > 100 * 1024 * 1024:
            raise ValueError("DOCX 解压体积超限")
        if len({item.filename for item in entries}) != len(entries):
            raise ValueError("DOCX 重复 ZIP 条目")
        texts = []
        for item in entries:
            name = item.filename
            if name.startswith(("word/media/", "word/embeddings/")) or name.endswith("vbaProject.bin"):
                raise ValueError("DOCX 包含未验证图片或嵌入对象，已阻断")
            if name.endswith(".rels"):
                relations = ET.fromstring(archive.read(name))
                if any(node.get("TargetMode") == "External" for node in relations):
                    raise ValueError("DOCX 外部链接未验证，已阻断")
            if name == "word/document.xml" or name.startswith(("word/header", "word/footer", "word/footnotes", "word/endnotes")):
                if not name.endswith(".xml"):
                    continue
                root = ET.fromstring(archive.read(name))
                blocked = {W + kind for kind in ("ins", "del", "moveFrom", "moveTo", "altChunk", "object", "drawing", "pict", "instrText", "fldSimple")}
                if any(node.tag in blocked for node in root.iter()):
                    raise ValueError("DOCX 包含修订、域或复杂对象，已阻断")
                for paragraph in root.iter(W + "p"):
                    texts.append("".join(node.text or "" for node in paragraph.iter(W + "t")))
        if not texts:
            raise ValueError("DOCX 无可验证正文")
        return "\n".join(texts)


def inspect_pdf(raw: bytes) -> str:
    with fitz.open(stream=raw, filetype="pdf") as document:
        if document.is_encrypted:
            raise ValueError("加密 PDF 未验证，已阻断")
        for page in document:
            for annotation in list(page.annots() or []):
                page.delete_annot(annotation)
        for page in document:
            if page.get_images(full=True) or page.get_drawings():
                raise ValueError("PDF 含未验证图片或矢量内容，已阻断安全文件导出")
            if not page.get_text().strip():
                raise ValueError("PDF 存在无文本页面，已阻断")
        with fitz.open(stream=document.tobytes(garbage=4), filetype="pdf") as cleaned:
            for xref in range(1, cleaned.xref_length()):
                obj = cleaned.xref_object(xref)
                if any(marker in obj for marker in ("/JavaScript", "/OpenAction", "/AA", "/AcroForm", "/RichMedia", "/XFA", "/OCProperties", "/Type/XObject", "/Type /XObject")):
                    raise ValueError("PDF 含未验证动态或复杂对象，已阻断")
        # Attachments, metadata, links and annotations are deliberately not copied.
        return "\n".join(page.get_text() for page in document)


def inspect_source(path: Path) -> str:
    raw = path.read_bytes()
    return inspect_pdf(raw) if path.suffix.lower() == ".pdf" else inspect_docx(raw)


def rebuild_safe_file(text: str, suffix: str) -> bytes:
    output = io.BytesIO()
    if suffix == ".docx":
        document = Document()
        properties = document.core_properties
        for key in ("author", "last_modified_by", "title", "subject", "keywords", "comments", "category"):
            setattr(properties, key, "")
        for line in text.splitlines():
            document.add_paragraph(line)
        document.save(output)
        return output.getvalue()
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    style = ParagraphStyle("safe", fontName="STSong-Light", fontSize=11, leading=16, wordWrap="CJK")
    document = SimpleDocTemplate(output, title="", author="", subject="")
    story = [Paragraph(escape(line) or "&#160;", style) for line in text.splitlines()]
    document.build(story)
    raw = output.getvalue()
    with fitz.open(stream=raw, filetype="pdf") as pdf:
        pdf.set_metadata({})
        pdf.del_xml_metadata()
        return pdf.tobytes(garbage=4, clean=True, deflate=True)


def verify_output(raw: bytes, suffix: str, original_values: list[str]) -> str:
    text = inspect_pdf(raw) if suffix == ".pdf" else inspect_docx(raw)
    if any(value and value in text for value in original_values):
        raise ValueError("安全副本仍包含已识别原值")
    if suffix == ".pdf":
        with fitz.open(stream=raw, filetype="pdf") as pdf:
            if pdf.embfile_count() or pdf.get_xml_metadata() or any(pdf.metadata.get(key) for key in ("author", "title", "subject", "keywords", "creator", "producer")):
                raise ValueError("PDF 附件或属性清理失败")
            if any(list(page.annots() or []) or list(page.widgets() or []) or page.get_links() for page in pdf):
                raise ValueError("PDF 批注、链接或表单清理失败")
    else:
        properties = Document(io.BytesIO(raw)).core_properties
        if any(getattr(properties, key) for key in ("author", "last_modified_by", "title", "subject", "keywords", "comments", "category")):
            raise ValueError("DOCX 文档属性清理失败")
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if any(name.startswith("word/comments") for name in archive.namelist()):
                raise ValueError("DOCX 批注清理失败")
            for name in archive.namelist():
                if name.endswith(".xml"):
                    root = ET.fromstring(archive.read(name))
                    content = " ".join(root.itertext())
                    if any(value and value in content for value in original_values):
                        raise ValueError("DOCX 隐藏 XML 仍含原值")
    return text
