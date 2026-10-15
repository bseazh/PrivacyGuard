import io
import tempfile
import unittest
import zipfile
from pathlib import Path

try:
    import fitz
    from docx import Document
    from privacyguard.safe_files import inspect_docx, inspect_pdf, rebuild_safe_file, verify_output
    from privacyguard.semantic import SemanticEnsemble
    from privacyguard.models import Finding
    from privacyguard.rules import merge_findings, replace_with_tokens
    from privacyguard.store import SafeReferenceStore, ReferenceError
    AVAILABLE = True
except ImportError:
    AVAILABLE = False


@unittest.skipUnless(AVAILABLE, "Safe file dependencies are installed in .venv")
class SafeFilesTests(unittest.TestCase):
    def test_overlapping_spans_cover_union(self):
        merged = merge_findings([Finding("PERSON", 0, 4, .9, "a"), Finding("ADDRESS", 2, 7, .8, "b")])
        self.assertEqual([(x.start, x.end) for x in merged], [(0, 7)])

    def test_repeated_values_reuse_tokens(self):
        text = "唐三与唐三"
        result = replace_with_tokens(text, [Finding("PERSON", 0, 2, .9, "a"), Finding("PERSON", 3, 5, .9, "b")])
        self.assertEqual(result, "<PERSON_001>与<PERSON_001>")

    def test_chunk_boundary_dedup_and_document_tail(self):
        class Redaction:
            def analyze_text(self, text):
                findings = []
                start = 0
                while (start := text.find("唐三", start)) >= 0:
                    findings.append({"type": "PERSON", "text": "唐三", "start": start, "end": start + 2, "confidence": None})
                    start += 2
                return findings
        text = "甲" * 2200 + "唐三" + "乙" * 8000 + "唐三"
        findings = SemanticEnsemble(Redaction()).analyze(text)
        self.assertEqual([(x.start, x.end) for x in findings], [(2200, 2202), (10202, 10204)])

    def test_docx_header_footer_and_clean_properties(self):
        document = Document()
        document.add_paragraph("正文")
        document.sections[0].header.paragraphs[0].text = "页眉 唐三"
        document.sections[0].footer.paragraphs[0].text = "页脚 alice@example.com"
        document.core_properties.author = "私人作者"
        stream = io.BytesIO()
        document.save(stream)
        text = inspect_docx(stream.getvalue())
        self.assertIn("页眉 唐三", text)
        self.assertIn("页脚 alice@example.com", text)
        raw = rebuild_safe_file(text.replace("唐三", "<PERSON_001>").replace("alice@example.com", "<EMAIL_001>"), ".docx")
        verify_output(raw, ".docx", ["唐三", "alice@example.com", "私人作者"])

    def test_docx_revisions_block(self):
        raw = rebuild_safe_file("正文", ".docx")
        target = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(target, "w") as output:
            for name in source.namelist():
                data = source.read(name)
                if name == "word/document.xml":
                    data = data.replace(b"<w:body>", b'<w:body><w:ins w:id="1"/>')
                output.writestr(name, data)
        with self.assertRaises(ValueError):
            inspect_docx(target.getvalue())

    def test_pdf_annotations_metadata_attachments_are_not_copied(self):
        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "Contact alice@example.com")
        page.add_text_annot((72, 100), "私人批注")
        pdf.set_metadata({"author": "私人作者"})
        pdf.embfile_add("secret.txt", b"attachment secret")
        text = inspect_pdf(pdf.tobytes())
        pdf.close()
        output = rebuild_safe_file(text.replace("alice@example.com", "<EMAIL_001>"), ".pdf")
        verify_output(output, ".pdf", ["alice@example.com", "私人作者", "私人批注"])

    def test_pdf_unverified_visual_content_blocks(self):
        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "Contact")
        page.draw_rect(fitz.Rect(72, 100, 200, 200))
        with self.assertRaises(ValueError):
            inspect_pdf(pdf.tobytes())
        pdf.close()

    def test_pdf_invisible_text_is_inspected(self):
        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "Visible text")
        page.insert_text((72, 100), "hidden@example.com", render_mode=3)
        self.assertIn("hidden@example.com", inspect_pdf(pdf.tobytes()))
        pdf.close()

    def test_output_known_original_values_block_in_both_formats(self):
        for suffix in (".pdf", ".docx"):
            with self.subTest(suffix=suffix), self.assertRaises(ValueError):
                verify_output(rebuild_safe_file("alice@example.com", suffix), suffix, ["alice@example.com"])

    def test_long_document_outputs_preserve_tail(self):
        text = "合同条款：付款期限需确认。\n" * 900 + "尾部 <PERSON_001>"
        for suffix in (".pdf", ".docx"):
            with self.subTest(suffix=suffix):
                extracted = verify_output(rebuild_safe_file(text, suffix), suffix, ["alice@example.com"])
                self.assertIn("<PERSON_001>", extracted)
                self.assertIn("尾部", extracted)

    def test_file_reference_detects_tampering(self):
        with tempfile.TemporaryDirectory() as directory:
            store = SafeReferenceStore(Path(directory), 60)
            path = store.root / "safe.txt"
            path.write_text("safe")
            reference = store.create(safe_text="safe", output_path=path, source_digest="x", counts={})
            path.write_text("changed")
            with self.assertRaises(ReferenceError):
                store.get(reference.token)


if __name__ == "__main__":
    unittest.main()
