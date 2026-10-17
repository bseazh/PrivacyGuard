from __future__ import annotations

import tempfile
import unittest
from unittest.mock import AsyncMock, patch
from pathlib import Path

from service.redaction_core import detect_entities, is_valid_china_id

try:
    from docx import Document
    from reportlab.pdfgen import canvas
    from privacyguard.config import Settings
    from privacyguard.core import PrivacyGuard
    from privacyguard.models import Finding
    from privacyguard.parsers import extract_local_text
    from privacyguard.security import SecurityError, mask_secret, validate_source_path
    from privacyguard.store import ReferenceError, SafeReferenceStore
    CORE_TEST_DEPENDENCIES = True
except ModuleNotFoundError:
    CORE_TEST_DEPENDENCIES = False


class FakePresidio:
    def analyze(self, text: str) -> list[Finding]:
        findings = []
        for needle, kind in (("alice@example.com", "EMAIL_ADDRESS"), ("13800138000", "PHONE_NUMBER")):
            start = text.find(needle)
            if start >= 0:
                findings.append(Finding(kind, start, start + len(needle), 0.9, "presidio"))
        return findings


@unittest.skipUnless(CORE_TEST_DEPENDENCIES, "PrivacyGuard Core test dependencies are installed in .venv")
class PrivacyGuardCoreTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name).resolve()
        self.settings = Settings(
            allowed_roots=(self.root,), data_dir=self.root / "state", max_file_bytes=1024 * 1024,
            reference_ttl_seconds=60, presidio_url="http://127.0.0.1:1", redaction_url="http://127.0.0.1:1", redaction_token=None,
        )
        self.core = PrivacyGuard(self.settings, presidio=FakePresidio())
        class NoSemantic:
            def analyze(self, text):
                return []
        self.core.semantic = NoSemantic()
        class FakeRedaction:
            def upload_and_parse(self, path):
                return {"file_id": "fixture", "content": extract_local_text(path), "entities": []}
        self.core.redaction = FakeRedaction()

    def tearDown(self):
        self.temp.cleanup()

    def test_scrubs_and_requires_safe_reference(self):
        original = "联系 alice@example.com 或 13800138000，身份证 11010519491231002X。"
        result = self.core.scrub_text(original)
        self.assertEqual(result["status"], "approved")
        self.assertNotIn("content", result)
        safe = self.core.get_safe_content(result["safe_reference"])
        self.assertEqual(safe["status"], "approved")
        self.assertNotIn("alice@example.com", safe["content"])
        self.assertNotIn("13800138000", safe["content"])
        self.assertNotIn("11010519491231002X", safe["content"])
        self.assertIn("联系", safe["content"])

    def test_scan_file_returns_reference_not_source_text(self):
        source = self.root / "contract.txt"
        source.write_text("邮箱 alice@example.com", encoding="utf-8")
        result = self.core.scan_file(str(source))
        self.assertEqual(result["status"], "approved")
        self.assertNotIn("alice@example.com", str(result))

    def test_redact_text_creates_independent_verified_file(self):
        source = self.root / "contract.txt"
        source.write_text("联系人 唐三，邮箱 alice@example.com", encoding="utf-8")
        result = self.core.redact_file(str(source))
        self.assertEqual(result["status"], "approved")
        self.assertEqual(result["output_kind"], "safe-file")
        safe = self.core.get_safe_content(result["safe_reference"])
        self.assertEqual(safe["kind"], "file")
        self.assertNotEqual(Path(safe["path"]), source)
        self.assertEqual(Path(safe["path"]).read_text(), safe["content"])
        self.assertNotIn("alice@example.com", safe["content"])

    def test_process_and_ask_hides_content_and_reference(self):
        source = self.root / "contract.txt"
        source.write_text("联系人 唐三，邮箱 alice@example.com", encoding="utf-8")
        self.core.secure_ask = AsyncMock(return_value={
            "status": "approved",
            "answer": "合同风险需要人工核查。",
            "task": "risk_summary",
            "guardrail": "privacyguard-presidio-pre-send",
        })
        result = __import__("asyncio").run(self.core.process_and_ask(str(source)))
        self.assertEqual(result["status"], "approved")
        self.assertNotIn("safe_reference", result)
        self.assertNotIn("content", result)
        self.assertIsNone(result["safe_file"])
        self.core.secure_ask.assert_awaited_once()

    def test_process_and_ask_blocks_without_model_fallback(self):
        result = __import__("asyncio").run(self.core.process_and_ask(str(self.root / "missing.txt")))
        self.assertEqual(result["status"], "blocked")
        self.assertEqual(result["stage"], "scan")
        self.assertNotIn("safe_reference", result)

    def test_txt_markdown_docx_and_text_pdf(self):
        sources = [
            self.root / "case.txt",
            self.root / "case.md",
            self.root / "case.markdown",
            self.root / "case.eml",
        ]
        for source in sources:
            source.write_text("联系 alice@example.com", encoding="utf-8")
        document = Document()
        document.add_paragraph("联系 alice@example.com")
        docx_path = self.root / "case.docx"
        document.save(docx_path)
        pdf_path = self.root / "case.pdf"
        pdf = canvas.Canvas(str(pdf_path))
        pdf.drawString(72, 760, "Contact alice@example.com")
        pdf.save()

        for source in (*sources, docx_path, pdf_path):
            with self.subTest(suffix=source.suffix):
                result = self.core.scan_file(str(source))
                self.assertEqual(result["status"], "approved")
                safe = self.core.get_safe_content(result["safe_reference"])
                self.assertNotIn("alice@example.com", safe["content"])

    def test_rejects_symlink_and_outside_root(self):
        source = self.root / "source.txt"
        source.write_text("safe", encoding="utf-8")
        link = self.root / "link.txt"
        link.symlink_to(source)
        with self.assertRaises(SecurityError):
            validate_source_path(link, (self.root,), 100)
        with self.assertRaises(SecurityError):
            validate_source_path("relative.txt", (self.root,), 100)

    def test_expired_reference_is_blocked(self):
        store = SafeReferenceStore(self.root / "short", ttl_seconds=-1)
        reference = store.create(safe_text="safe", output_path=None, source_digest="x", counts={})
        with self.assertRaises(ReferenceError):
            store.get(reference.token)

    def test_presidio_failure_is_fail_closed(self):
        class Broken:
            def analyze(self, text):
                from privacyguard.adapters.presidio import PresidioUnavailable
                raise PresidioUnavailable("offline")

        core = PrivacyGuard(self.settings, presidio=Broken())
        result = core.scrub_text("alice@example.com")
        self.assertEqual(result["status"], "blocked")
        self.assertNotIn("alice@example.com", str(result))

    def test_recheck_new_entity_is_added_to_original_list(self):
        class Semantic:
            def analyze(self, text):
                if "alice@example.com" not in text:
                    start = text.find("史莱克学院")
                    if start >= 0:
                        return [Finding("ORGANIZATION", start, start + 5, .8, "late")]
                return []
        self.core.semantic = Semantic()
        result = self.core.scrub_text("alice@example.com 联系史莱克学院")
        self.assertEqual(result["status"], "approved")
        self.assertEqual(result["counts"]["ORGANIZATION"], 1)
        self.assertNotIn("史莱克学院", self.core.get_safe_content(result["safe_reference"])["content"])

    def test_known_names_expand_to_all_occurrences(self):
        class Semantic:
            def analyze(self, text):
                start = text.find("唐三")
                return [Finding("PERSON", start, start + 2, .9, "first-only")] if start >= 0 else []
        self.core.semantic = Semantic()
        result = self.core.scrub_text("唐三交付，唐三复核。")
        self.assertEqual(result["counts"], {"PERSON": 2})
        safe = self.core.get_safe_content(result["safe_reference"])["content"]
        self.assertEqual(safe.count("<PERSON_001>"), 2)

    def test_malformed_docx_returns_blocked_without_source(self):
        path = self.root / "broken.docx"
        path.write_bytes(b"alice@example.com")
        result = self.core.scan_file(str(path))
        self.assertEqual(result["status"], "blocked")
        self.assertNotIn("alice@example.com", str(result))

    def test_failed_export_does_not_publish_reference(self):
        document = Document()
        document.add_paragraph("alice@example.com")
        path = self.root / "contract.docx"
        document.save(path)
        with patch("privacyguard.core.verify_output", side_effect=ValueError("private: alice@example.com")):
            result = self.core.redact_file(str(path))
        self.assertEqual(result["status"], "blocked")
        self.assertIsNone(result["safe_reference"])
        self.assertNotIn("alice@example.com", str(result))
        self.assertEqual(list(self.core.store.root.iterdir()), [])

    def test_failed_upstream_never_falls_back_to_local_text(self):
        source = self.root / "contract.txt"
        source.write_text("邮箱 alice@example.com")
        for payload in ({"content": "", "entities": []},
                        {"content": "邮箱 alice@example.com", "entities": [
                            {"start": 0, "end": 2, "text": "wrong", "type": "PERSON"}]}):
            with patch.object(self.core.redaction, "upload_and_parse", return_value=payload):
                result = self.core.scan_file(str(source))
            self.assertEqual(result["status"], "blocked")
            self.assertIsNone(result["safe_reference"])
        self.assertEqual(list(self.core.store.root.iterdir()), [])

    def test_analysis_calls_upstream_without_rebuilding_file(self):
        source = self.root / "contract.txt"
        source.write_text("邮箱 alice@example.com")
        with patch.object(self.core.redaction, "upload_and_parse", wraps=self.core.redaction.upload_and_parse) as parse, \
                patch("privacyguard.core.rebuild_safe_file", side_effect=AssertionError("analysis must not export")):
            result = self.core.scan_file(str(source))
        self.assertEqual(result["status"], "approved")
        parse.assert_called_once_with(source)
        self.assertEqual(self.core.get_safe_content(result["safe_reference"])["kind"], "text")
        self.assertFalse(list(self.core.store.root.glob("*.txt")))

    def test_ocr_page_coverage_must_match_complete_body(self):
        source = self.root / "scan.pdf"
        pdf = canvas.Canvas(str(source))
        pdf.drawString(72, 760, "Page one")
        pdf.save()
        for pages, page_count in ((["Page one", "Page two"], 1), (["Page one", ""], 2),
                                  (["Page one", "Page two"], 2)):
            with self.subTest(pages=pages), patch.object(self.core.redaction, "upload_and_parse", return_value={
                "content": "Page one", "entities": [], "ocr_complete": True,
                "page_count": page_count, "pages": pages,
            }):
                result = self.core.scan_file(str(source))
                self.assertEqual(result["status"], "blocked")
                self.assertIsNone(result["safe_reference"])

    def test_docx_headers_omitted_by_upstream_are_still_protected(self):
        document = Document()
        document.add_paragraph("合同正文")
        document.sections[0].header.paragraphs[0].text = "邮箱 alice@example.com"
        document.sections[0].footer.paragraphs[0].text = "电话 13800138000"
        source = self.root / "headers.docx"
        document.save(source)
        with patch.object(self.core.redaction, "upload_and_parse", return_value={"content": "合同正文", "entities": []}):
            result = self.core.redact_file(str(source))
        self.assertEqual(result["status"], "approved", result)
        content = self.core.get_safe_content(result["safe_reference"])["content"]
        self.assertNotIn("alice@example.com", content)
        self.assertNotIn("13800138000", content)
        self.assertIn("邮箱", content)
        self.assertIn("电话", content)

    def test_source_mutation_during_parse_blocks_reference(self):
        source = self.root / "contract.txt"
        source.write_text("合同正文")
        def mutate(path):
            path.write_text("已改变")
            return {"content": "合同正文", "entities": []}
        with patch.object(self.core.redaction, "upload_and_parse", side_effect=mutate):
            self.assertEqual(self.core.scan_file(str(source))["status"], "blocked")
        self.assertEqual(list(self.core.store.root.iterdir()), [])

    def test_partial_pdf_text_cannot_hide_an_empty_page(self):
        source = self.root / "mixed.pdf"
        pdf = canvas.Canvas(str(source))
        pdf.drawString(72, 760, "Contract text")
        pdf.showPage()
        pdf.showPage()
        pdf.save()
        with patch.object(self.core.redaction, "upload_and_parse", return_value={"content": "Contract text", "entities": []}):
            self.assertEqual(self.core.scan_file(str(source))["status"], "blocked")
        self.assertEqual(list(self.core.store.root.iterdir()), [])

    def test_low_confidence_semantic_entity_blocks_safe_export(self):
        document = Document()
        document.add_paragraph("项目名称 星斗计划")
        path = self.root / "uncertain.docx"
        document.save(path)

        class UncertainSemantic:
            def analyze(self, text):
                start = text.find("星斗计划")
                return [Finding("PROJECT", start, start + 4, .75, "ner")] if start >= 0 else []

        self.core.semantic = UncertainSemantic()
        result = self.core.redact_file(str(path))
        self.assertEqual(result["status"], "blocked")
        self.assertIn("需要人工确认", result["reason"])
        self.assertIsNone(result["safe_reference"])

    def test_accepts_verified_upstream_entity_with_none_confidence(self):
        source = self.root / "case.txt"
        source.write_text("当事人唐三提交合同")

        class Redaction:
            def upload_and_parse(self, path):
                return {
                    "content": "当事人唐三提交合同",
                    "is_scanned": False,
                    "entities": [{"type": "PERSON", "text": "唐三", "start": 3, "end": 5, "confidence": None}],
                }

        self.core.redaction = Redaction()
        result = self.core.scan_file(str(source))
        self.assertEqual(result["status"], "approved")
        safe = self.core.get_safe_content(result["safe_reference"])
        self.assertNotIn("唐三", safe["content"])
        self.assertEqual(result["counts"], {"PERSON": 1})


class ChinaIdTests(unittest.TestCase):
    def test_validates_date_and_checksum(self):
        self.assertTrue(is_valid_china_id("11010519491231002X"))
        self.assertFalse(is_valid_china_id("11010519491331002X"))
        self.assertFalse(is_valid_china_id("110105194912310021"))
        text = "真 11010519491231002X 假 11010519491331002X"
        values = [item.value for item in detect_entities(text) if item.type == "ID_NUMBER"]
        self.assertEqual(values, ["11010519491231002X"])

    def test_reference_mask_never_returns_complete_secret(self):
        token = "abcd0123456789wxyz"
        masked = mask_secret(token)
        self.assertEqual(masked, "abcd...wxyz")
        self.assertNotIn(token, masked)


if __name__ == "__main__":
    unittest.main()
